# -*- coding: utf-8 -*-
"""Audio-gid uchun o'qish matnlari (docs/05_Audio_gid_matnlari.docx).
Ishga tushirish: python tools/build_audio_script.py   (node.js kerak — matnlar narrations.js dan olinadi)"""
import json, os, subprocess
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
js = """
const fs=require('fs'); const src=f=>fs.readFileSync(f,'utf8');
const ctx=new Function(src('assets/js/data.js')+src('assets/js/narrations.js')+';return {PLACES,REGIONS,NARR};')();
console.log(JSON.stringify(ctx));"""
data = json.loads(subprocess.check_output(["node", "-e", js], cwd=ROOT))
places, regions, narr = data["PLACES"], {r["id"]: r for r in data["REGIONS"]}, data["NARR"]

d = Document()
st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(13)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
h = d.add_heading("Tourist.uz — audio-gid uchun o'qish matnlari (o'zbek tili)", 0)
p = d.add_paragraph()
p.add_run("Qanday yozish kerak:\n").bold = True
p.add_run(
    "1. Telefon diktofonida jim xonada, telefonni og'izdan 20–30 sm uzoqlikda ushlab yozing.\n"
    "2. Matnni shoshilmasdan, aniq va iliq ohangda o'qing (1–1,5 daqiqa).\n"
    "3. Faylni pastda ko'rsatilgan nom bilan saqlang (masalan: registan.mp3 yoki registan.m4a).\n"
    "4. Barcha fayllarni GitHub'da assets/audio/uz/ papkasiga yuklang — qolganini dasturchi ulaydi "
    "(assets/audio/manifest.js).\n"
    "Matnni o'zgartirmoqchi bo'lsangiz, avval assets/js/narrations.js dagi matnni ham shunga moslang — ekranda o'sha matn ko'rsatiladi.")
for r_id in [r for r in regions]:
    items = [pl for pl in places if pl["region"] == r_id]
    if not items: continue
    d.add_heading(regions[r_id]["name"]["uz"], 1)
    for pl in items:
        d.add_heading(f'{pl["name"]["uz"]}', 2)
        meta = d.add_paragraph()
        r = meta.add_run(f'Fayl nomi: {pl["id"]}.mp3'); r.bold = True; r.font.color.rgb = RGBColor(0x1D, 0x3F, 0x8F)
        text = narr.get(pl["id"], {}).get("uz") or pl["desc"]["uz"]
        for para in text.split("\n\n"):
            d.add_paragraph(para)
out = os.path.join(ROOT, "docs", "05_Audio_gid_matnlari.docx")
d.save(out)
print("✓", os.path.relpath(out, ROOT), len(places), "obida")
